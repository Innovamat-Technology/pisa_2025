"""Render Catalan, Spanish and English LATAM reports and historical Q4-Q1 gaps.

Reads existing numerical CSV outputs; no estimation or recalibration is performed.
Run after 14_latam_quartiles.py and 15_latam_history.py. Both the 2022–2025
comparison and the historical reports are produced in ca, es and en.
"""
from __future__ import annotations

import argparse
import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

from lib import config as C

ROOT = C.RESULTS / "latam"
COUNTRIES = ["ARG", "BRA", "CHL", "COL", "MEX", "PER", "URY"]
YEARS = [2015, 2018, 2022, 2025]
NAMES = {
    "ca": dict(ARG="Argentina", BRA="Brasil", CHL="Xile", COL="Colòmbia", MEX="Mèxic", PER="Perú", URY="Uruguai", LATAM7="Mitjana simple"),
    "es": dict(ARG="Argentina", BRA="Brasil", CHL="Chile", COL="Colombia", MEX="México", PER="Perú", URY="Uruguay", LATAM7="Media simple"),
    "en": dict(ARG="Argentina", BRA="Brazil", CHL="Chile", COL="Colombia", MEX="Mexico", PER="Peru", URY="Uruguay", LATAM7="Unweighted country average"),
}
LABELS = {
    "ca": {"escs":"ESCS publicat", "escs_h":"ESCS harmonitzat", "homepos":"HOMEPOS publicat", "homepos_h":"HOMEPOS harmonitzat", "hisei_h":"Ocupació parental (HISEI)", "pared_h":"Educació parental (PARED)", "books_h":"Llibres a casa", "escs_pca":"ESCS, primer component principal", "escs_cc":"ESCS publicat, casos complets", "escs_h_cc":"ESCS harmonitzat, casos complets"},
    "es": {"escs":"ESCS publicado", "escs_h":"ESCS armonizado", "homepos":"HOMEPOS publicado", "homepos_h":"HOMEPOS armonizado", "hisei_h":"Ocupación parental (HISEI)", "pared_h":"Educación parental (PARED)", "books_h":"Libros en casa", "escs_pca":"ESCS, primer componente principal", "escs_cc":"ESCS publicado, casos completos", "escs_h_cc":"ESCS armonizado, casos completos"},
    "en": {"escs":"Published ESCS", "escs_h":"Harmonized ESCS", "homepos":"Published HOMEPOS", "homepos_h":"Harmonized HOMEPOS", "hisei_h":"Parental occupation (HISEI)", "pared_h":"Parental education (PARED)", "books_h":"Books at home", "escs_pca":"ESCS, first principal component", "escs_cc":"Published ESCS, complete cases", "escs_h_cc":"Harmonized ESCS, complete cases"},
}
UI = {
    "ca": dict(country="País", index="Índex", year="Any", gap="Bretxa Q4−Q1", points="Punts PISA · IC del 95%", change="Canvi en punts PISA · IC del 95%", current_title="Matemàtiques per quartil socioeconòmic a l’Amèrica Llatina", history_title="Sèries històriques de matemàtiques per quartil, 2015–2025", current_chart="Matemàtiques: canvi 2022–2025 per quartil d’{index}", levels_chart="Matemàtiques per quartil · {index}", history_chart="Matemàtiques per quartil · {index} · 2015–2025", notes="Notes", history_notes="Quartils de l’índex publicat de cada cicle. Argentina: sèrie des del 2018.", current_notes="Q1: quartil inferior · Q4: superior. Harmonització recalibrada amb els set països.", qlabels=["Q1 · inferior", "Q2", "Q3", "Q4 · superior"]),
    "es": dict(country="País", index="Índice", year="Año", gap="Brecha Q4−Q1", points="Puntos PISA · IC del 95%", change="Cambio en puntos PISA · IC del 95%", current_title="Matemáticas por cuartil socioeconómico en América Latina", history_title="Series históricas de matemáticas por cuartil, 2015–2025", current_chart="Matemáticas: cambio 2022–2025 por cuartil de {index}", levels_chart="Matemáticas por cuartil · {index}", history_chart="Matemáticas por cuartil · {index} · 2015–2025", notes="Notas", history_notes="Cuartiles del índice publicado en cada ciclo. Argentina: serie desde 2018.", current_notes="Q1: cuartil inferior · Q4: superior. Armonización recalibrada con los siete países.", qlabels=["Q1 · inferior", "Q2", "Q3", "Q4 · superior"]),
    "en": dict(country="Country", index="Index", year="Year", gap="Q4−Q1 gap", points="PISA score points · 95% CI", change="Change in PISA score points · 95% CI", current_title="Mathematics by socioeconomic quartile in Latin America", history_title="Historical mathematics scores by quartile, 2015–2025", current_chart="Mathematics: 2022–2025 change by {index} quartile", levels_chart="Mathematics by quartile · {index}", history_chart="Mathematics by quartile · {index} · 2015–2025", notes="Notes", history_notes="Quartiles of the published index in each cycle. Argentina: series starts in 2018.", current_notes="Q1: bottom quartile · Q4: top quartile. Harmonization recalibrated on the seven countries.", qlabels=["Q1 · bottom", "Q2", "Q3", "Q4 · top"]),
}
for _lang in ["ca", "es", "en"]:
    LABELS[_lang]["escs_h4"] = LABELS[_lang]["escs_h"]
    LABELS[_lang]["homepos_h4"] = LABELS[_lang]["homepos_h"]
UI["ca"].update(history_title="Evolució de la bretxa Q4−Q1 en matemàtiques, 2015–2025",
                history_chart="Bretxa Q4−Q1 en matemàtiques · {index} · 2015–2025",
                gap_axis="Bretxa Q4−Q1 · punts PISA · IC del 95%",
                history_notes="Índex publicat i harmonització amb 13 ítems comuns als quatre cicles. Argentina: des del 2018.")
UI["es"].update(history_title="Evolución de la brecha Q4−Q1 en matemáticas, 2015–2025",
                history_chart="Brecha Q4−Q1 en matemáticas · {index} · 2015–2025",
                gap_axis="Brecha Q4−Q1 · puntos PISA · IC del 95%",
                history_notes="Índice publicado y armonización con 13 ítems comunes a los cuatro ciclos. Argentina: desde 2018.")
UI["en"].update(history_title="Q4−Q1 mathematics gap over time, 2015–2025",
                history_chart="Q4−Q1 mathematics gap · {index} · 2015–2025",
                gap_axis="Q4−Q1 gap · PISA score points · 95% CI",
                history_notes="Published index and harmonization using 13 items common to all four cycles. Argentina: from 2018.")


def phrase(lang, es, en, ca=None):
    return {"es":es, "en":en, "ca":ca if ca is not None else es}[lang]


def number(value, lang, signed=False, digits=1):
    if not np.isfinite(value):
        return "—"
    text = format(value, f"{ '+' if signed else ''}.{digits}f").replace("-", "−")
    return text if lang == "en" else text.replace(".", ",")


def interval(value, se, lang):
    return f"{number(value,lang,True)} [{number(value-1.96*se,lang,True)}; {number(value+1.96*se,lang,True)}]"


def table(headers, rows):
    return ["| " + " | ".join(headers) + " |", "|" + "|".join(["---"]*len(headers)) + "|"] + ["| " + " | ".join(map(str,row)) + " |" for row in rows]


def load_tables(folder):
    return {p.stem:pd.read_csv(p) for p in (folder / "tables").glob("*.csv")}


def file_hash(path):
    with Path(path).open("rb") as f:
        return hashlib.file_digest(f, "sha256").hexdigest()


def write_manifest(out, source, lang):
    data = {"generated_utc":datetime.now(timezone.utc).isoformat(), "language":lang,
            "rendering_script_sha256":file_hash(__file__), "numerical_inputs":{
                p.name:file_hash(p) for p in (source / "tables").glob("*.csv")},
            "estimates_recomputed":False}
    (out / "render_manifest.json").write_text(json.dumps(data,indent=2) + "\n")


def plot_style():
    for p in (C.ROOT / "assets/fonts").glob("*.ttf"):
        font_manager.fontManager.addfont(str(p))
    plt.rcParams.update({"font.family":"Source Sans 3", "font.size":11, "pdf.fonttype":42,
                         "axes.spines.top":False, "axes.spines.right":False, "legend.frameon":False})


def save(fig, folder, name):
    folder.mkdir(parents=True, exist_ok=True)
    for ext in ["png", "pdf"]:
        fig.savefig(folder / f"{name}.{ext}", dpi=220, bbox_inches="tight")
    plt.close(fig)


def panels(lang):
    fig, axes = plt.subplots(3,3,figsize=(12,9.8),sharex=True,sharey=True,squeeze=False)
    for ax, cnt in zip(axes.flat, COUNTRIES):
        ax.set_title(NAMES[lang][cnt],loc="left",fontweight="semibold")
        ax.grid(axis="y",color=".9",lw=.6)
        ax.tick_params(axis="x",labelbottom=True)
    for ax in list(axes.flat)[7:]:
        ax.set_visible(False)
    return fig, axes


def current_figures(tables, out, lang):
    u = UI[lang]
    colors = ["#2b5c8a", "#C0370C"]
    for family in ["escs", "homepos"]:
        fig, axes = panels(lang)
        for ax, cnt in zip(axes.flat, COUNTRIES):
            for j, idx in enumerate([family, family+"_h"]):
                t = tables["quartile_changes_math"]
                g = t[(t.cnt == cnt) & (t["index"] == idx) & t.group.isin(["Q1","Q2","Q3","Q4"])].sort_values("group")
                assert len(g) == 4
                ax.errorbar(np.arange(1,5)+(j-.5)*.05,g.estimate,yerr=1.96*g.se,color=colors[j],marker="o",lw=1.6,capsize=3,label=LABELS[lang][idx])
            ax.axhline(0,color=".45",lw=.8)
            ax.set_xticks(range(1,5),["Q1","Q2","Q3","Q4"])
        handles, labels = axes.flat[0].get_legend_handles_labels()
        fig.legend(handles,labels,loc="lower center",ncol=2,bbox_to_anchor=(.5,.015))
        fig.suptitle(u["current_chart"].format(index=family.upper()),fontsize=16,x=.07,ha="left")
        fig.supylabel(u["change"],x=.015)
        fig.text(.5,-.008,u["current_notes"],ha="center",fontsize=10)
        fig.tight_layout(rect=(.035,.08,1,.96))
        save(fig,out/"figures",f"quartile_changes_{family}_math")
    for idx in ["escs", "escs_h"]:
        fig, axes = panels(lang)
        for ax, cnt in zip(axes.flat, COUNTRIES):
            for j,year in enumerate([2022,2025]):
                t = tables["quartile_levels_math"]
                g = t[(t.cnt == cnt) & (t["index"] == idx) & (t.wave == year) & t.group.isin(["Q1","Q2","Q3","Q4"])].sort_values("group")
                assert len(g) == 4
                ax.errorbar(np.arange(1,5)+(j-.5)*.05,g.estimate,yerr=1.96*g.se,color=colors[j],marker="o",lw=1.6,capsize=3,label=str(year))
            ax.set_xticks(range(1,5),["Q1","Q2","Q3","Q4"])
        handles,labels = axes.flat[0].get_legend_handles_labels()
        fig.legend(handles,labels,loc="lower center",ncol=2,bbox_to_anchor=(.5,.015))
        fig.suptitle(u["levels_chart"].format(index=LABELS[lang][idx]),fontsize=16,x=.07,ha="left")
        fig.supylabel(u["points"],x=.015)
        fig.tight_layout(rect=(.035,.08,1,.96))
        save(fig,out/"figures",f"quartile_levels_{idx}_math")


def header(name, lang):
    labels = {
        "cnt":("Código ISO3","ISO3 code","Codi ISO3"), "country":("País","Country","País"),
        "wave":("Año","Year","Any"), "index":("Índice","Index","Índex"), "group":("Grupo","Group","Grup"),
        "quartile":("Cuartil","Quartile","Quartil"), "estimate":("Estimación","Estimate","Estimació"),
        "se":("Error estándar","Standard error","Error estàndard"), "ci_low":("IC 95% inferior","95% CI lower","IC 95% inferior"),
        "ci_high":("IC 95% superior","95% CI upper","IC 95% superior"), "n":("Alumnos","Students","Alumnes"),
        "n_total":("Alumnos totales","Total students","Alumnes totals"), "n_valid":("Alumnos válidos","Valid students","Alumnes vàlids"),
        "n_raw":("Alumnos en el fichero","Students in source file","Alumnes al fitxer"),
        "n_systems":("Número de países","Number of countries","Nombre de països"), "countries":("Países incluidos","Included countries","Països inclosos"),
        "status":("Estado","Status","Estat"), "included":("Incluido","Included","Inclòs"), "in_trend":("Incluido en tendencia","Included in trend","Inclòs a la tendència"),
        "valid_weight_pct":("Peso válido (%)","Valid weight (%)","Pes vàlid (%)"), "weight_pct":("Peso del cuartil (%)","Quartile weight (%)","Pes del quartil (%)"),
        "link_se":("Error de enlace","Linking error","Error d’enllaç"), "published":("Índice publicado","Published index","Índex publicat"),
        "harmonised":("Índice armonizado","Harmonized index","Índex harmonitzat"), "source":("Fuente","Source","Font"),
        "computed":("Calculado","Computed","Calculat"), "official":("OCDE publicado","OECD published","OCDE publicat"),
        "difference":("Diferencia de validación","Validation difference","Diferència de validació"), "statistic":("Estadístico","Statistic","Estadístic"),
        "source_table":("Tabla de referencia","Reference table","Taula de referència"), "field":("Campo","Field","Camp"),
        "new":("Nuevo cálculo","New calculation","Nou càlcul"), "existing":("Cálculo anterior","Existing calculation","Càlcul anterior"),
        "item":("Ítem","Item","Ítem"), "a":("Discriminación IRT","IRT discrimination","Discriminació IRT"), "steps":("Umbrales IRT","IRT thresholds","Llindars IRT"),
    }
    if str(name) in labels:
        return phrase(lang,*labels[str(name)])
    replacements = {"change":phrase(lang,"cambio","change","canvi"), "se":phrase(lang,"error estándar","standard error","error estàndard"),
                    "mean":phrase(lang,"media","mean","mitjana"), "weight":phrase(lang,"peso (%)","weight (%)","pes (%)"),
                    "n":phrase(lang,"alumnos","students","alumnes"), "estimate":phrase(lang,"estimación","estimate","estimació"),
                    "history":phrase(lang,"serie histórica","historical series","sèrie històrica"),
                    "existing":phrase(lang,"cálculo anterior","existing calculation","càlcul anterior")}
    return " · ".join(replacements.get(x,x) for x in str(name).split("_"))


def localized_table(frame, lang):
    t = frame.copy()
    if "country" in t and "cnt" in t:
        t["country"] = t.cnt.map(NAMES[lang])
    for column in ["index","published","harmonised"]:
        if column in t:
            t[column] = t[column].map(lambda x:LABELS[lang].get(x,x))
    statuses = {
        "both cycles":phrase(lang,"Ambos ciclos","Both cycles","Tots dos cicles"),
        "estimated":phrase(lang,"Estimado","Estimated","Estimat"),
        "available":phrase(lang,"Disponible","Available","Disponible"),
        "not_comparable_sampling":phrase(lang,"No comparable: cobertura de la muestra","Not comparable: sampling coverage","No comparable: cobertura de la mostra"),
        "absent_from_file":phrase(lang,"Ausente del fichero","Absent from source file","Absent del fitxer"),
    }
    if "status" in t:
        t["status"] = t.status.map(lambda x:statuses.get(x,x))
    if "wave" in t:
        t["wave"] = t.wave.replace({"change":phrase(lang,"Cambio","Change","Canvi")})
    for column in ["statistic","field"]:
        if column in t:
            t[column] = t[column].map(lambda x:header(x,lang))
    t.columns = [header(x,lang) for x in t.columns]
    return t


def workbook(tables, lines, out, lang, name):
    sheet_names = {
        "quartile_levels_math":("Niveles","Levels","Nivells"), "quartile_changes_math":("Cambios","Changes","Canvis"),
        "quartiles_math":("Resumen cuartiles","Quartile summary","Resum quartils"), "coverage":("Cobertura","Coverage","Cobertura"),
        "quartile_check_math":("Pesos cuartiles","Quartile weights","Pesos quartils"), "quartile_checks":("Pesos cuartiles","Quartile weights","Pesos quartils"),
        "overall_math":("Medias nacionales","National means","Mitjanes nacionals"), "harmonisation_effect_math":("Efecto armonización","Harmonization effect","Efecte harmonització"),
        "irt_parameters":("Parámetros IRT","IRT parameters","Paràmetres IRT"), "existing_python_comparison":("Validación Python","Python validation","Validació Python"),
        "official_comparison":("Referencia OCDE","OECD reference","Referència OCDE"), "gap_summary_math":("Brechas","Gaps","Bretxes"),
        "country_availability":("Disponibilidad","Availability","Disponibilitat"), "availability":("Disponibilidad","Availability","Disponibilitat"),
        "levels_math":("Niveles por año","Levels by year","Nivells per any"), "levels_wide_math":("Series","Series","Sèries"),
        "gaps_math":("Brechas por año","Gaps by year","Bretxes per any"), "gap_series_math":("Series de brechas","Gap series","Sèries de bretxes"),
        "latest_validation":("Validación 2022 2025","2022 2025 validation","Validació 2022 2025"),
        "latest_gap_validation":("Validación brechas","Gap validation","Validació bretxes"),
    }
    dictionary = []
    with pd.ExcelWriter(out/name,engine="openpyxl") as writer:
        pd.DataFrame({UI[lang]["notes"]:[line for line in lines if line and not line.startswith(("|","![","#"))]}).to_excel(writer,sheet_name=UI[lang]["notes"],index=False)
        for key, data in tables.items():
            sheet = phrase(lang,*sheet_names[key])
            localized_table(data,lang).to_excel(writer,sheet_name=sheet,index=False)
            dictionary.extend([{"CSV":key+".csv", "Sheet":sheet, "Field":str(col), "Label":header(col,lang)} for col in data.columns])
        dictionary_name = phrase(lang,"Diccionario","Dictionary","Diccionari")
        pd.DataFrame(dictionary).rename(columns={"Sheet":phrase(lang,"Hoja","Sheet","Full"), "Field":phrase(lang,"Campo original","Original field","Camp original"), "Label":phrase(lang,"Etiqueta","Label","Etiqueta")}).to_excel(writer,sheet_name=dictionary_name,index=False)
        for sheet in writer.book:
            sheet.freeze_panes="A2"
            sheet.auto_filter.ref=sheet.dimensions
            for cell in sheet[1]:
                cell.font=Font(bold=True,color="FFFFFF")
                cell.fill=PatternFill("solid",fgColor="2B5C8A")
            for column in sheet.columns:
                width=min(45,max(14,max(len(str(c.value or "")) for c in list(column)[:60])+2))
                sheet.column_dimensions[column[0].column_letter].width=width
            if sheet.title==UI[lang]["notes"]:
                sheet.column_dimensions["A"].width=115
                for row in sheet.iter_rows(min_row=2):
                    row[0].alignment=Alignment(wrap_text=True,vertical="top")
                    sheet.row_dimensions[row[0].row].height=max(30,15*(len(str(row[0].value))//105+1))
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value,float):
                        cell.number_format="0.00"


def current_report(tables, validation, out, lang):
    u=UI[lang]
    q=tables["quartiles_math"].set_index(["cnt","index"])
    p=lambda es,en,ca:phrase(lang,es,en,ca)
    workbook_name="latam_mathematiques.xlsx" if lang=="ca" else "latam_mathematics.xlsx"
    source_prefix="" if lang=="ca" else "../"
    lines=["# "+u["current_title"],"",p("Comparación PISA 2022–2025. Países: ","PISA 2022–2025 comparison. Countries: ","Comparació PISA 2022–2025. Països: ")+", ".join(NAMES[lang][c] for c in COUNTRIES)+".","",
           p("Los siete países tienen datos en ambos ciclos y entran en la calibración conjunta. Esta selección no comprende todos los participantes de la región.","All seven countries have data in both cycles and enter the pooled calibration. This selection does not comprise all participants in the region.","Els set països tenen dades als dos cicles i entren en la calibració conjunta. Aquesta selecció no inclou tots els participants de la regió."),"",
           "## "+p("Resultados principales","Main findings","Resultats principals"),""]
    for pub,har,name in [("escs","escs_h","ESCS"),("homepos","homepos_h","HOMEPOS")]:
        a,b=q.loc["LATAM7",pub],q.loc["LATAM7",har]
        lines.append(f"- **{name}:** "+p("cambio medio de la brecha Q4−Q1 de ","average change in the Q4−Q1 gap of ","canvi mitjà de la bretxa Q4−Q1 de ")+interval(a["Q4-Q1_change"],a["Q4-Q1_se"],lang)+p(" puntos con el índice publicado y "," points using the published index and "," punts amb l’índex publicat i ")+interval(b["Q4-Q1_change"],b["Q4-Q1_se"],lang)+p(" con el armonizado (IC del 95%)."," using the harmonized index (95% CI)."," amb l’harmonitzat (IC del 95%)."))
    r=q.loc["LATAM7","escs_h"]
    lines.append(p("- Con el ESCS armonizado, Q1 cambia ","- With harmonized ESCS, Q1 changes by ","- Amb l’ESCS harmonitzat, Q1 canvia ")+number(r.Q1_change,lang,True)+p(" puntos y Q4, "," points and Q4 by "," punts i Q4, ")+number(r.Q4_change,lang,True)+p(" puntos."," points."," punts."))
    sig=[NAMES[lang][c] for c in COUNTRIES if abs(q.loc[c,"escs_h"]["Q4-Q1_change"])>1.96*q.loc[c,"escs_h"]["Q4-Q1_se"]]
    lines.append(p("- Países cuyo cambio de brecha armonizada tiene un IC del 95% que excluye el cero: ","- Countries whose harmonized gap change has a 95% CI excluding zero: ","- Països amb un canvi de bretxa harmonitzada amb un IC del 95% que exclou el zero: ")+", ".join(sig)+".")
    low=tables["coverage"].query("index == 'escs_h'").sort_values("valid_weight_pct").iloc[0]
    lines.append(p("- La menor cobertura del ESCS armonizado corresponde a ","- The lowest harmonized ESCS coverage is in ","- La menor cobertura de l’ESCS harmonitzat correspon a ")+NAMES[lang][low.cnt]+f", {int(low.wave)}: "+number(low.valid_weight_pct,lang)+p("% del peso de la muestra. Los casos completos permiten comparar los índices sobre los mismos alumnos, pero no corrigen posibles sesgos de no respuesta.","% of sample weight. Complete-case estimates compare the indices on the same students, but do not correct potential non-response bias.","% del pes de la mostra. Els casos complets permeten comparar els índexs sobre els mateixos alumnes, però no corregeixen possibles biaixos de no-resposta."))
    lines += ["","## "+p("Cambio de la brecha Q4−Q1","Change in the Q4−Q1 gap","Canvi de la bretxa Q4−Q1"),"",p("Puntos PISA; IC del 95% entre corchetes. Un valor negativo indica una reducción de la brecha.","PISA score points; 95% CIs in brackets. Negative values indicate a narrowing gap.","Punts PISA; IC del 95% entre claudàtors. Un valor negatiu indica una reducció de la bretxa."),""]
    rows=[]
    for c in COUNTRIES+["LATAM7"]:
        rows.append([NAMES[lang][c]]+[interval(q.loc[c,i]["Q4-Q1_change"],q.loc[c,i]["Q4-Q1_se"],lang) for i in ["escs","escs_h"]])
    lines += table([u["country"],LABELS[lang]["escs"],LABELS[lang]["escs_h"]],rows)
    lines += ["",p("La media otorga el mismo peso a cada país.","The average gives each country equal weight.","La mitjana dona el mateix pes a cada país."),"","## "+p("Cambios por cuartil","Changes by quartile","Canvis per quartil"),""]
    rows=[]
    for c in COUNTRIES+["LATAM7"]:
        for idx in ["escs","escs_h","homepos","homepos_h"]:
            rows.append([NAMES[lang][c],LABELS[lang][idx]]+[number(q.loc[c,idx][f"Q{k}_change"],lang,True) for k in range(1,5)])
    lines += table([u["country"],u["index"],"Q1","Q2","Q3","Q4"],rows)
    lines += ["","## "+p("Niveles de matemáticas en 2025","Mathematics score levels in 2025","Nivells de matemàtiques el 2025"),"",p("Puntos PISA. Los errores estándar e intervalos se incluyen en el Excel y en la tabla de niveles.","PISA score points. Standard errors and confidence intervals are included in the workbook and the score-level table.","Punts PISA. Els errors estàndard i els intervals de confiança s’inclouen a l’Excel i a la taula de nivells."),""]
    rows=[]
    for c in COUNTRIES:
        for idx in ["escs","escs_h"]:
            r=q.loc[c,idx]
            rows.append([NAMES[lang][c],LABELS[lang][idx]]+[number(r[f"Q{k}_2025"],lang) for k in range(1,5)]+[number(r["Q4-Q1_2025"],lang)])
    lines += table([u["country"],u["index"],"Q1","Q2","Q3","Q4",u["gap"]],rows)
    lines += ["","## "+p("Gráficos","Figures","Gràfics"),"","![ESCS](figures/quartile_changes_escs_math.png)","","![ESCS](figures/quartile_levels_escs_h_math.png)","","## "+p("Método e interpretación","Method and interpretation","Mètode i interpretació"),"",
        p("Se han extraído los alumnos de los ficheros originales, incluidos los países no miembros de la OCDE. Los cuartiles se definen dentro de cada país y año, con aproximadamente el 25% del peso final en cada grupo y desempate aleatorio con semilla 7. Q1 y Q4 son posiciones relativas nacionales, no niveles de renta idénticos entre países. Los ciclos son muestras distintas, no un seguimiento de los mismos alumnos.","Students were extracted from the original files, including non-OECD countries. Quartiles are defined within each country and year, with approximately 25% of final student weight in each group and random tie-breaking with seed 7. Q1 and Q4 are national relative positions, not identical income levels across countries. Cycles are separate samples, not a panel of the same students.","S’han extret els alumnes dels fitxers originals, incloent-hi els països no membres de l’OCDE. Els quartils es defineixen dins de cada país i any, amb aproximadament el 25% del pes final per grup i desempat aleatori amb llavor 7. Q1 i Q4 són posicions relatives nacionals, no nivells de renda idèntics entre països. Els cicles són mostres diferents, no un seguiment dels mateixos alumnes."),"",
        p("Se utilizan diez valores plausibles y 80 réplicas BRR con Fay 0,5, recalculando los cuartiles en cada réplica. La media internacional suma las matrices de covarianza nacionales independientes y divide por el número de países al cuadrado, como en la especificación principal en R. El error de enlace de 1,220 puntos se añade una sola vez a los cambios de medias, también a la media internacional; se cancela en los cambios de brecha. Los intervalos son estimación ± 1,96 errores estándar.","Estimates use ten plausible values and 80 BRR replicates with Fay 0.5, recomputing quartiles in every replicate. The international mean sums independent national covariance matrices and divides by the squared number of countries, as in the main R specification. The 1.220-point linking error is added once to changes in means, including the international mean; it cancels in gap changes. Intervals are estimate ± 1.96 standard errors.","S’utilitzen deu valors plausibles i 80 rèpliques BRR amb Fay 0,5, recalculant els quartils a cada rèplica. La mitjana internacional suma les matrius de covariància nacionals independents i divideix pel nombre de països al quadrat, com a l’especificació principal en R. L’error d’enllaç d’1,220 punts s’afegeix una sola vegada als canvis de mitjanes, també a la mitjana internacional; es cancel·la en els canvis de bretxa. Els intervals són estimació ± 1,96 errors estàndard."),"",
        p("HOMEPOS armonizado utiliza los mismos 16 ítems comunes, recodificaciones y funciones Python del proyecto: modelo 2PL/GPCM, calibración conjunta 2022+2025 con igual peso por país-año y puntuaciones WLE con al menos diez respuestas. ESCS combina HISEI, PARED (3→6 y 14,5→14) y HOMEPOS común. Un único componente ausente se imputa por regresión dentro del país-año con un residuo aleatorio (semilla 20252022); los componentes se estandarizan conjuntamente y se promedian con igual peso.","Harmonized HOMEPOS uses the same 16 common items, recodes and Python functions as the project: a 2PL/GPCM model, pooled 2022+2025 calibration with equal country-year weight, and WLE scores requiring at least ten answers. ESCS combines HISEI, PARED (3→6 and 14.5→14) and common-item HOMEPOS. A single missing component is imputed by within-country-year regression with a random residual (seed 20252022); components are standardized jointly and averaged with equal weights.","HOMEPOS harmonitzat utilitza els mateixos 16 ítems comuns, recodificacions i funcions Python del projecte: model 2PL/GPCM, calibració conjunta 2022+2025 amb el mateix pes per país-any i puntuacions WLE amb almenys deu respostes. ESCS combina HISEI, PARED (3→6 i 14,5→14) i HOMEPOS comú. Un únic component absent s’imputa per regressió dins del país-any amb un residu aleatori (llavor 20252022); els components s’estandarditzen conjuntament i se’n fa la mitjana amb pesos iguals."),"",
        p("**La calibración y la estandarización se han rehecho con los siete países.** Esta extensión regional no conserva exactamente los resultados de la calibración OCDE del repositorio y no constituye una escala oficial de tendencia de la OCDE. Los errores estándar son condicionales a los índices reconstruidos: no se recalibra el IRT ni se repite la imputación dentro de cada réplica.","**Calibration and standardization have been refitted on the seven countries.** This regional extension does not exactly preserve the repository’s OECD-calibrated results and is not an official OECD trend scale. Standard errors are conditional on the reconstructed indices: IRT is not recalibrated and imputation is not repeated within each replicate.","**La calibració i l’estandardització s’han refet amb els set països.** Aquesta extensió regional no conserva exactament els resultats de la calibració OCDE del repositori i no constitueix una escala oficial de tendència de l’OCDE. Els errors estàndard són condicionals als índexs reconstruïts: no es recalibra l’IRT ni es repeteix la imputació dins de cada rèplica."),"",
        p("También se incluyen HISEI, PARED, libros en casa, el primer componente principal y comparaciones de ESCS sobre los mismos casos completos. La tabla del efecto de armonización conserva la covarianza entre índices. Los intervalos no están ajustados por comparaciones múltiples y los resultados descriptivos no identifican efectos causales.","The workbook also includes HISEI, PARED, books at home, the first principal component and ESCS comparisons on the same complete cases. The harmonization-effect table retains covariance between indices. Intervals are not adjusted for multiple comparisons, and descriptive results do not identify causal effects.","També s’inclouen HISEI, PARED, llibres a casa, el primer component principal i comparacions d’ESCS sobre els mateixos casos complets. La taula de l’efecte d’harmonització conserva la covariància entre índexs. Els intervals no estan ajustats per comparacions múltiples i els resultats descriptius no identifiquen efectes causals."),"",
        "## "+p("Validación","Validation","Validació"),"",
        p("Identificadores, pesos y valores plausibles alineados: comprobados. Desviación máxima respecto al 25% de peso por cuartil: ","Identifiers, weights and plausible-value alignment: checked. Maximum deviation from 25% weight per quartile: ","Identificadors, pesos i valors plausibles alineats: comprovats. Desviació màxima respecte del 25% de pes per quartil: ")+number(validation["max_quartile_weight_deviation_pp"],lang,digits=3)+p(" puntos porcentuales."," percentage points."," punts percentuals."),"",
        f"{validation['existing_comparison_cells']} "+p("celdas contrastadas con el cálculo Python original; discrepancia máxima ","cells checked against the original Python calculation; maximum discrepancy ","cel·les contrastades amb el càlcul Python original; discrepància màxima ")+f"{validation['existing_max_abs_difference']:.3g}. "+str(validation["aggregate_checks"])+p(" comprobaciones de las medias internacionales y sus errores estándar."," checks of international means and their standard errors."," comprovacions de les mitjanes internacionals i els seus errors estàndard."),"",
        p("Diferencia máxima frente a las tablas OCDE I.B1.2b.24 e I.B1.2b.38: ","Maximum difference from OECD Tables I.B1.2b.24 and I.B1.2b.38: ","Diferència màxima respecte de les taules OCDE I.B1.2b.24 i I.B1.2b.38: ")+number(validation["official_max_abs_estimate_difference"],lang,digits=3)+p(" puntos en las estimaciones y "," score points in estimates and "," punts en les estimacions i ")+number(validation["official_max_abs_se_difference"],lang,digits=3)+p(" en los errores estándar. Los empates, los procedimientos de varianza y las revisiones de los ficheros públicos pueden generar discrepancias; se documentan sin forzar coincidencias."," in standard errors. Ties, variance procedures and public-file revisions can cause discrepancies; these are documented without forcing agreement."," en els errors estàndard. Els empats, els procediments de variància i les revisions dels fitxers públics poden generar discrepàncies; es documenten sense forçar coincidències."),"",
        "## "+p("Archivos y reproducción","Files and reproduction","Fitxers i reproducció"),"",
        f"- [Excel]({workbook_name})",
        f"- [CSV]({source_prefix}tables/)",
        "- "+p("[Versiones e informes](../index.md)","[Report versions](../index.md)","[Versions i informes](index.md)"),
        "- "+p("[Series históricas de brechas](../history/es/README.md)","[Historical gap series](../history/en/README.md)","[Sèries històriques de bretxes](history/README.md)"),
        "- "+p("[Fuentes y configuración](../sources.json)","[Sources and configuration](../sources.json)","[Fonts i configuració](sources.json)"),"",
        "```bash",".venv/bin/python code/python/scripts/14_latam_quartiles.py",".venv/bin/python code/python/scripts/16_latam_reports.py --current-only","```","",
        p("El primer comando requiere los originales de 2022 y 2025 en `data/raw/` o `PISA_RAW_DIR`. El segundo traduce y representa las mismas tablas, sin recalcular estimaciones. `--refresh` en el primer comando rehace la extracción y calibración; las cachés están en `data/interim/python/latam/`.","The first command requires the 2022 and 2025 originals in `data/raw/` or `PISA_RAW_DIR`. The second translates and renders the same tables without re-estimating them. `--refresh` on the first command rebuilds extraction and calibration; caches are under `data/interim/python/latam/`.","El primer comandament requereix els originals del 2022 i del 2025 a `data/raw/` o `PISA_RAW_DIR`. El segon tradueix i representa les mateixes taules, sense recalcular estimacions. `--refresh` al primer comandament refà l’extracció i la calibració; les memòries cau són a `data/interim/python/latam/`."),"",
        "## "+p("Fuentes","Sources","Fonts"),"",
        p("Microdatos oficiales, copias locales documentadas en `data/README.md`: ","Official microdata, using local copies documented in `data/README.md`: ","Microdades oficials, còpies locals documentades a `data/README.md`: ")+"[PISA 2022](https://www.oecd.org/en/data/datasets/pisa-2022-database.html), [PISA 2025](https://www.oecd.org/en/data/datasets/pisa-2025-database.html). "+p("Descargados el 18 de septiembre de 2026. Portales consultados el 23 de septiembre de 2026.","Downloaded on 18 September 2026. Portals consulted on 23 September 2026.","Descarregades el 18 de setembre de 2026. Portals consultats el 23 de setembre de 2026."),"",
        "[OCDE / OECD](https://stat.link/k68msa). "+p("Las constantes de enlace y los documentos técnicos están descritos en `reference/README.md`.","Linking constants and technical documents are described in `reference/README.md`.","Les constants d’enllaç i els documents tècnics es descriuen a `reference/README.md`."),""]
    (out/"README.md").write_text("\n".join(lines))
    workbook(tables,lines,out,lang,workbook_name)


def history_figures(tables,out,lang):
    t=tables["gaps_math"]
    specifications = [
        ("fig_gap_series_math", ["escs","escs_h4","homepos","homepos_h4"], "ESCS / HOMEPOS"),
        ("fig_gap_series_escs_math", ["escs","escs_h4"], "ESCS"),
        ("fig_gap_series_homepos_math", ["homepos","homepos_h4"], "HOMEPOS"),
    ]
    for filename,indices,family in specifications:
        fig,axes=panels(lang)
        for ax,cnt in zip(axes.flat,COUNTRIES):
            for idx in indices:
                g=t[(t.cnt == cnt) & (t["index"] == idx)].sort_values("wave")
                assert len(g)==(3 if cnt=="ARG" else 4)
                ax.errorbar(g.wave,g.estimate,yerr=1.96*g.se,
                            color="#C0370C" if idx.endswith("_h4") else "#2b5c8a",
                            linestyle="--" if idx.startswith("homepos") else "-",
                            marker="s" if idx.startswith("homepos") else "o",
                            ms=3.5,lw=1.5,elinewidth=.65,capsize=2,label=LABELS[lang][idx])
            ax.set_xticks(YEARS,[str(y) for y in YEARS])
            ax.set_xlim(2014.5,2025.5)
            if cnt=="ARG":
                ax.text(.02,.95,phrase(lang,"Desde 2018","From 2018","Des del 2018"),transform=ax.transAxes,va="top",fontsize=9,color=".4")
        handles,labels=axes.flat[0].get_legend_handles_labels()
        fig.legend(handles,labels,loc="lower center",ncol=2,bbox_to_anchor=(.5,.015))
        fig.suptitle(UI[lang]["history_chart"].format(index=family),fontsize=16,x=.07,ha="left")
        fig.supylabel(UI[lang]["gap_axis"],x=.015)
        fig.text(.5,-.008,UI[lang]["history_notes"],ha="center",fontsize=10)
        fig.tight_layout(rect=(.035,.11,1,.96))
        save(fig,out/"figures",filename)


def history_report(tables,validation,out,lang):
    p=lambda es,en,ca:phrase(lang,es,en,ca)
    u=UI[lang]
    source="https://www.oecd.org/en/publications/pisa-2018-results-volume-i_5f07c754-en/full-report/component-9.html"
    diagnostics=json.loads((ROOT/"history/sources.json").read_text())["harmonization"]
    lines=["# "+u["history_title"],"",", ".join(NAMES[lang][c] for c in COUNTRIES)+".","",
        p("Se muestra la **brecha Q4−Q1 de matemáticas en cada ciclo**, con el índice publicado y nuestro índice armonizado, tanto para ESCS como para HOMEPOS. Cada punto es la puntuación media del cuartil superior menos la del inferior. Se representa el nivel de esa brecha, en puntos PISA, sin calcular su cambio entre años, siguiendo `fig_gap_series_math.png` del proyecto.","The series show the **Q4−Q1 mathematics gap in each cycle**, using the published index and our harmonized index, for both ESCS and HOMEPOS. Each point is the mean score in the top quartile minus the mean in the bottom quartile. Charts plot the level of that gap in PISA score points, without calculating changes between years, following the project's `fig_gap_series_math.png`.","Es mostra la **bretxa Q4−Q1 de matemàtiques a cada cicle**, amb l’índex publicat i el nostre índex harmonitzat, tant per a ESCS com per a HOMEPOS. Cada punt és la puntuació mitjana del quartil superior menys la de l’inferior. Es representa el nivell d’aquesta bretxa, en punts PISA, sense calcular-ne el canvi entre anys, seguint `fig_gap_series_math.png` del projecte."),"",
        "## "+p("Cobertura temporal","Time coverage","Cobertura temporal"),""]
    rows=[[NAMES[lang][c],"2018, 2022, 2025" if c=="ARG" else "2015, 2018, 2022, 2025"] for c in COUNTRIES]
    lines += table([u["country"],p("Ciclos incluidos","Included cycles","Cicles inclosos")],rows)
    lines += ["",p("**Argentina empieza en 2018.** La OCDE considera que los resultados nacionales de 2015 no son comparables debido a un marco muestral incompleto. No se sustituye Argentina por la Ciudad Autónoma de Buenos Aires ni se interpola 2015.","**Argentina starts in 2018.** OECD considers the 2015 national results non-comparable because of an incomplete sampling frame. Argentina is not replaced by the City of Buenos Aires, and 2015 is not interpolated.","**Argentina comença el 2018.** L’OCDE considera que els resultats nacionals del 2015 no són comparables per un marc mostral incomplet. No se substitueix Argentina per la Ciutat Autònoma de Buenos Aires ni s’interpola el 2015.")+f" [OCDE / OECD]({source}).","",
              "## "+p("Gráficos","Figures","Gràfics"),"","![ESCS / HOMEPOS](figures/fig_gap_series_math.png)","",
              p("Azul: índice publicado. Naranja: armonizado. Línea continua: ESCS. Línea discontinua: HOMEPOS. Barras: intervalos de confianza del 95%.","Blue: published index. Orange: harmonized index. Solid line: ESCS. Dashed line: HOMEPOS. Bars: 95% confidence intervals.","Blau: índex publicat. Taronja: harmonitzat. Línia contínua: ESCS. Línia discontínua: HOMEPOS. Barres: intervals de confiança del 95%."),"",
              p("Versiones separadas: [ESCS](figures/fig_gap_series_escs_math.png) y [HOMEPOS](figures/fig_gap_series_homepos_math.png). Todos los gráficos están disponibles también en PDF.","Separate charts: [ESCS](figures/fig_gap_series_escs_math.png) and [HOMEPOS](figures/fig_gap_series_homepos_math.png). All charts are also available as PDFs.","Versions separades: [ESCS](figures/fig_gap_series_escs_math.png) i [HOMEPOS](figures/fig_gap_series_homepos_math.png). Tots els gràfics també estan disponibles en PDF."),""]
    for family in ["escs","homepos"]:
        lines += ["## "+family.upper(),"",p("Brecha Q4−Q1 en puntos PISA. Los errores estándar y los intervalos del 95% están en el Excel y en los CSV.","Q4−Q1 gap in PISA score points. Standard errors and 95% intervals are in the workbook and CSV tables.","Bretxa Q4−Q1 en punts PISA. Els errors estàndard i els intervals del 95% són a l’Excel i als CSV."),""]
        rows=[]
        for c in COUNTRIES:
            for idx in [family,family+"_h4"]:
                g=tables["gaps_math"]
                g=g[(g.cnt==c)&(g["index"]==idx)].set_index("wave")
                rows.append([NAMES[lang][c],LABELS[lang][idx]]+
                            [number(g.loc[year,"estimate"],lang) if year in g.index else "—" for year in YEARS])
        lines += table([u["country"],u["index"]]+[str(y) for y in YEARS],rows)+[""]
    lines += ["## "+p("Método e interpretación","Method and interpretation","Mètode i interpretació"),"",
        p("Se utilizan los diez valores plausibles de matemáticas y los pesos finales oficiales. Cada cuartil reúne aproximadamente el 25% del peso de los alumnos con índice y valores plausibles válidos en su país y ciclo. Los cuartiles se recalculan con cada una de las 80 réplicas BRR (Fay 0,5); los empates se resuelven con semilla 7. Los errores estándar combinan varianza muestral y varianza entre valores plausibles, conservando la covarianza entre Q1 y Q4 al calcular la brecha. Se requieren al menos 200 alumnos válidos por índice y país-año.","Estimates use all ten mathematics plausible values and official final weights. Each quartile contains approximately 25% of the weight of students with a valid index and plausible values in its country and cycle. Quartiles are recomputed for each of the 80 BRR replicates (Fay 0.5); ties are broken using seed 7. Standard errors combine sampling and between-plausible-value variance, retaining the covariance between Q1 and Q4 when calculating the gap. At least 200 valid students are required for each index and country-year.","S’utilitzen els deu valors plausibles de matemàtiques i els pesos finals oficials. Cada quartil reuneix aproximadament el 25% del pes dels alumnes amb índex i valors plausibles vàlids al seu país i cicle. Els quartils es recalculen amb cadascuna de les 80 rèpliques BRR (Fay 0,5); els empats es resolen amb llavor 7. Els errors estàndard combinen variància mostral i variància entre valors plausibles, conservant la covariància entre Q1 i Q4 en calcular la bretxa. Calen almenys 200 alumnes vàlids per índex i país-any."),"",
        p("**Índice publicado:** ESCS y HOMEPOS tal como aparecen en el fichero original de cada ciclo, sin aplicar reescalados retrospectivos. Su definición puede cambiar con el cuestionario.","**Published index:** ESCS and HOMEPOS as released in each cycle’s original file, with no retrospective rescaling. Their definitions may change with the questionnaire.","**Índex publicat:** ESCS i HOMEPOS tal com apareixen al fitxer original de cada cicle, sense aplicar reescalats retrospectius. La seva definició pot canviar amb el qüestionari."),"",
        p("**Nuestro índice armonizado:** HOMEPOS se reconstruye con los 13 ítems comunes a los cuatro ciclos, recodificados a categorías comparables, mediante un modelo 2PL/GPCM con parámetros únicos. La calibración conjunta da el mismo peso total a cada una de las 27 combinaciones país-año disponibles; las puntuaciones WLE requieren al menos ocho respuestas. En el ítem de ordenadores se conserva como ausente la combinación «ninguno + respuesta desconocida», siguiendo la corrección de la implementación R del proyecto.","**Our harmonized index:** HOMEPOS is rebuilt from the 13 items common to all four cycles, recoded to comparable categories, using one set of 2PL/GPCM parameters. Pooled calibration gives equal total weight to each of the 27 available country-year cells; WLE scores require at least eight answers. For the computer item, a combination of “none + unknown response” remains missing, following the correction in the project's R implementation.","**El nostre índex harmonitzat:** HOMEPOS es reconstrueix amb els 13 ítems comuns als quatre cicles, recodificats a categories comparables, mitjançant un model 2PL/GPCM amb paràmetres únics. La calibració conjunta dona el mateix pes total a cadascuna de les 27 combinacions país-any disponibles; les puntuacions WLE requereixen almenys vuit respostes. A l’ítem d’ordinadors es conserva com a absent la combinació «cap + resposta desconeguda», seguint la correcció de la implementació R del projecte."),"",
        p("ESCS combina HISEI, PARED y el HOMEPOS común. La educación parental de 2015 se reconstruye desde HISCED con la correspondencia modal observada en 2018, y PARED se recodifica con 3→6 y 14,5→14 años. Un único componente ausente se imputa por regresión dentro del país-año, añadiendo un residuo aleatorio con semilla 20252022. Los tres componentes se estandarizan en el conjunto de los cuatro ciclos, con igual peso total por país-año, se promedian con igual peso y el resultado se vuelve a estandarizar.","ESCS combines HISEI, PARED and common-item HOMEPOS. Parental education in 2015 is reconstructed from HISCED using the modal mapping observed in 2018, and PARED is recoded with 3→6 and 14.5→14 years. A single missing component is imputed by within-country-year regression, adding a random residual with seed 20252022. The three components are standardized across the four-cycle pool with equal country-year weight, averaged with equal component weights, and the result is standardized again.","ESCS combina HISEI, PARED i l’HOMEPOS comú. L’educació parental del 2015 es reconstrueix des d’HISCED amb la correspondència modal observada el 2018, i PARED es recodifica amb 3→6 i 14,5→14 anys. Un únic component absent s’imputa per regressió dins del país-any, afegint-hi un residu aleatori amb llavor 20252022. Els tres components s’estandarditzen en el conjunt dels quatre cicles, amb el mateix pes total per país-any, se’n fa la mitjana amb pesos iguals i el resultat es torna a estandarditzar."),"",
        p("**La armonización histórica utiliza 13 ítems y los cuatro ciclos; la comparación 2022–2025 utiliza 16 ítems y dos ciclos.** Por ello, los resultados armonizados de 2022 y 2025 pueden diferir entre ambos informes. Toda la curva histórica se estima con una única especificación común. La reconstrucción es regional y no constituye una escala oficial de tendencia de la OCDE. Los errores estándar son condicionales a los índices reconstruidos: el modelo IRT y la imputación no se vuelven a estimar en las réplicas.","**Historical harmonization uses 13 items and all four cycles; the 2022–2025 comparison uses 16 items and two cycles.** Harmonized results for 2022 and 2025 can therefore differ between the two reports. The entire historical curve uses one common specification. This is a regional reconstruction, not an official OECD trend scale. Standard errors are conditional on the reconstructed indices: IRT and imputation are not re-estimated in the replicates.","**L’harmonització històrica utilitza 13 ítems i els quatre cicles; la comparació 2022–2025 utilitza 16 ítems i dos cicles.** Per això, els resultats harmonitzats del 2022 i del 2025 poden diferir entre tots dos informes. Tota la corba històrica s’estima amb una única especificació comuna. La reconstrucció és regional i no constitueix una escala oficial de tendència de l’OCDE. Els errors estàndard són condicionals als índexs reconstruïts: el model IRT i la imputació no es tornen a estimar a les rèpliques."),"",
        p("Los intervalos son estimación ± 1,96 errores estándar. El error aditivo de enlace de la escala de rendimiento se cancela en el contraste Q4−Q1 del mismo ciclo. El solapamiento visual de intervalos no es una prueba del cambio entre años. Q1 y Q4 representan posiciones relativas dentro de cada país; la serie no sigue a los mismos alumnos ni a un grupo socioeconómico fijo, y no identifica efectos causales. Se presentan resultados nacionales; no se calcula una media regional con una composición que cambie entre ciclos.","Intervals are estimate ± 1.96 standard errors. The additive achievement-scale linking error cancels in the within-cycle Q4−Q1 contrast. Visual overlap of intervals is not a test of change between years. Q1 and Q4 represent relative positions within each country; the series does not follow the same students or a fixed socioeconomic group and does not identify causal effects. Results are national; no regional average with changing country coverage is calculated.","Els intervals són estimació ± 1,96 errors estàndard. L’error additiu d’enllaç de l’escala de rendiment es cancel·la en el contrast Q4−Q1 del mateix cicle. El solapament visual dels intervals no és una prova del canvi entre anys. Q1 i Q4 representen posicions relatives dins de cada país; la sèrie no segueix els mateixos alumnes ni un grup socioeconòmic fix, i no identifica efectes causals. Es presenten resultats nacionals; no es calcula una mitjana regional amb una composició que canviï entre cicles."),"",
        "## "+p("Validación","Validation","Validació"),"",
        p("La serie contiene ","The series contains ","La sèrie conté ")+str(validation["country_cycles"])+p(" combinaciones país-año, "," country-year cells, "," combinacions país-any, ")+str(validation["gap_rows"])+p(" brechas y "," gaps and "," bretxes i ")+str(validation["level_rows"])+p(" estimaciones por cuartil e índice. Identificadores, pesos, valores plausibles, tamaños de grupo, medias ponderadas y Q4−Q1: comprobados."," quartile/index estimates. Identifiers, weights, plausible values, group sizes, weighted means and Q4−Q1: checked."," estimacions per quartil i índex. Identificadors, pesos, valors plausibles, mides de grup, mitjanes ponderades i Q4−Q1: comprovats."),"",
        p("El modelo IRT convergió en ","The IRT model converged in ","El model IRT va convergir en ")+str(diagnostics["irt_iterations"])+p(" iteraciones, con "," iterations, with "," iteracions, amb ")+str(diagnostics["calibration_students"])+p(" alumnos en el conjunto de calibración. Los parámetros completos están en `irt_parameters.csv` y la configuración en `sources.json`."," students in the calibration pool. Full parameters are in `irt_parameters.csv` and settings in `sources.json`."," alumnes al conjunt de calibració. Els paràmetres complets són a `irt_parameters.csv` i la configuració a `sources.json`."),"",
        str(validation["matched_recent_level_rows"])+p(" niveles de los índices publicados en 2022 y 2025 reproducen el análisis anterior; discrepancia máxima: "," score levels based on the published indices in 2022 and 2025 reproduce the previous analysis; maximum discrepancy: "," nivells dels índexs publicats del 2022 i del 2025 reprodueixen l’anàlisi anterior; discrepància màxima: ")+f"{validation['recent_max_abs_estimate_difference']:.3g}"+p(" puntos."," points."," punts."),"",
        p("La comparación con las tablas OCDE I.B1.2b.24 e I.B1.2b.38 se conserva en `official_comparison.csv`. En 2015 y 2018, esas tablas retrospectivas pueden utilizar índices reescalados: no son un objetivo de igualdad para los cuartiles del índice original. Las diferencias se documentan por país y ciclo.","Comparison with OECD Tables I.B1.2b.24 and I.B1.2b.38 is retained in `official_comparison.csv`. For 2015 and 2018, these retrospective tables may use rescaled indices, so they are not an equality target for quartiles of the original index. Differences are documented by country and cycle.","La comparació amb les taules OCDE I.B1.2b.24 i I.B1.2b.38 es conserva a `official_comparison.csv`. El 2015 i el 2018, aquestes taules retrospectives poden utilitzar índexs reescalats: no són un objectiu d’igualtat per als quartils de l’índex original. Les diferències es documenten per país i cicle."),"",
        "## "+p("Archivos y reproducción","Files and reproduction","Fitxers i reproducció"),"",
        "- [Excel](historical_mathematics.xlsx)",
        "- [CSV]("+("tables/" if lang=="ca" else "../tables/")+")",
        "- "+p("[Versiones e informes](../../index.md)","[Report versions](../../index.md)","[Versions i informes](../index.md)"),
        "- "+p("[Fuentes y huellas SHA-256](../sources.json)","[Sources and SHA-256 fingerprints](../sources.json)","[Fonts i empremtes SHA-256](sources.json)"),"",
        "```bash",".venv/bin/python code/python/scripts/15_latam_history.py",".venv/bin/python code/python/scripts/16_latam_reports.py","```","",
        p("Se requieren los originales de 2015, 2018, 2022 y 2025 en `data/raw/` o `PISA_RAW_DIR`. Los extractos y cachés se guardan en `data/interim/python/latam/history/`. Las versiones lingüísticas representan exactamente las mismas tablas numéricas.","The 2015, 2018, 2022 and 2025 originals must be in `data/raw/` or `PISA_RAW_DIR`. Extracts and caches are stored under `data/interim/python/latam/history/`. Language versions render exactly the same numerical tables.","Calen els originals del 2015, 2018, 2022 i 2025 a `data/raw/` o `PISA_RAW_DIR`. Els extractes i memòries cau es guarden a `data/interim/python/latam/history/`. Les versions lingüístiques representen exactament les mateixes taules numèriques."),"",
        "## "+p("Fuentes","Sources","Fonts"),"",
        ", ".join(f"[PISA {year}](https://www.oecd.org/en/data/datasets/pisa-{year}-database.html)" for year in YEARS)+".","",
        p("Se han reutilizado los ficheros oficiales disponibles localmente, documentados en `data/README.md`. Los portales y la advertencia sobre Argentina 2015 se consultaron el 23 de septiembre de 2026.","The locally available official files documented in `data/README.md` were reused. Portals and the Argentina 2015 comparability note were checked on 23 September 2026.","S’han reutilitzat els fitxers oficials disponibles localment, documentats a `data/README.md`. Els portals i l’advertiment sobre Argentina 2015 es van consultar el 23 de setembre de 2026."),""]
    (out/"README.md").write_text("\n".join(lines))
    workbook(tables,lines,out,lang,"historical_mathematics.xlsx")


def landing_page():
    lines=["# PISA · Amèrica Llatina / América Latina / Latin America","",
           "Argentina · Brasil / Brazil · Xile / Chile · Colòmbia / Colombia · Mèxic / México / Mexico · Perú / Peru · Uruguai / Uruguay","",
           "| Anàlisi / Análisis / Analysis | Català | Castellano | English |",
           "|---|---|---|---|",
           "| 2022–2025 | [Informe](README.md) · [Excel](latam_mathematiques.xlsx) | [Informe](es/README.md) · [Excel](es/latam_mathematics.xlsx) | [Report](en/README.md) · [Excel](en/latam_mathematics.xlsx) |"]
    if (ROOT/"history/en/README.md").exists():
        lines += ["| 2015–2025 · Q4−Q1 | [Informe](history/README.md) · [Excel](history/historical_mathematics.xlsx) | [Informe](history/es/README.md) · [Excel](history/es/historical_mathematics.xlsx) | [Report](history/en/README.md) · [Excel](history/en/historical_mathematics.xlsx) |","",
                  "| Gràfic / Gráfico / Chart · Q4−Q1 | Català | Castellano | English |","|---|---|---|---|",
                  "| ESCS + HOMEPOS | [PNG](history/figures/fig_gap_series_math.png) · [PDF](history/figures/fig_gap_series_math.pdf) | [PNG](history/es/figures/fig_gap_series_math.png) · [PDF](history/es/figures/fig_gap_series_math.pdf) | [PNG](history/en/figures/fig_gap_series_math.png) · [PDF](history/en/figures/fig_gap_series_math.pdf) |","",
                  "Argentina: 2018, 2022, 2025. Brasil / Brazil, Chile, Colombia, México / Mexico, Perú / Peru, Uruguay: 2015, 2018, 2022, 2025.","",
                  "Historical charts compare published and harmonized indices in each cycle. Their common calibration uses 13 items across four cycles; the 2022–2025 analysis uses 16 items across two cycles."]
    (ROOT/"index.md").write_text("\n".join(lines)+"\n")


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--current-only",action="store_true")
    args=parser.parse_args()
    plot_style()
    current=load_tables(ROOT)
    validation=json.loads((ROOT/"validation.json").read_text())
    for lang in ["ca","es","en"]:
        out=ROOT if lang=="ca" else ROOT/lang
        out.mkdir(parents=True,exist_ok=True)
        current_figures(current,out,lang)
        current_report(current,validation,out,lang)
        write_manifest(out,ROOT,lang)
        print(f"Rendered 2022–2025 report, workbook and figures: {lang}",flush=True)
    if not args.current_only:
        base=ROOT/"history"
        historical=load_tables(base)
        if not historical:
            raise FileNotFoundError("Run 15_latam_history.py before rendering historical reports")
        validation=json.loads((base/"validation.json").read_text())
        for lang in ["ca","es","en"]:
            out=base if lang=="ca" else base/lang
            out.mkdir(parents=True,exist_ok=True)
            history_figures(historical,out,lang)
            history_report(historical,validation,out,lang)
            write_manifest(out,base,lang)
            print(f"Rendered historical report, workbook and figures: {lang}",flush=True)
    landing_page()


if __name__=="__main__":
    main()
